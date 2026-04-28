"""Test case management routes."""



from fastapi import APIRouter, Body, Depends, File, Query, UploadFile, status
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.dependencies import get_current_user, get_db, success_response
from app.core.exceptions import ValidationError
from app.services.test_case_service import TestCaseService
from app.utils.helpers import build_filters

router = APIRouter(prefix="/testcases", tags=["Test Cases"])


# ── Bulk operation schemas ─────────────────────────────────────────────────


class _BulkIds(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=500, description="Test case IDs")


class _BulkApprovalTransition(_BulkIds):
    status: str = Field(min_length=1, max_length=30)


class _BulkUpdate(_BulkIds):
    """Plain-field bulk edit. Status (approval) goes through its own
    endpoint because of the workflow rules. ``unassign=true`` forces
    assignee → null; otherwise omit ``assigned_to`` to leave it
    untouched."""
    priority: str | None = Field(default=None, max_length=30)
    type_: str | None = Field(default=None, max_length=50, alias="type")
    automation_status: str | None = Field(default=None, max_length=30)
    assigned_to: str | None = Field(default=None, description="Target user ID")
    unassign: bool = Field(default=False, description="Force-clear the assignee")

    model_config = {"populate_by_name": True}


@router.get(
    "",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="List test cases",
)
async def list_test_cases(
    project_id: str | None = Query(default=None, description="Filter by project ID"),
    module_id: str | None = Query(default=None, description="Filter by module ID"),
    epic_id: str | None = Query(default=None, description="Filter by epic ID"),
    user_story_id: str | None = Query(default=None, description="Filter by user story ID"),
    test_status: str | None = Query(default=None, alias="status", description="Filter by status"),
    test_type: str | None = Query(default=None, alias="type", description="Filter by type"),
    priority: str | None = Query(default=None, description="Filter by priority"),
    automation_status: str | None = Query(default=None, description="Filter by automation status"),
    search: str | None = Query(default=None, description="Search by title or custom ID"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=10000),
    sort_by: str = Query(default="created_at"),
    sort_order: str = Query(default="desc", pattern="^(asc|desc)$"),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """List test cases with pagination, filtering, and search capabilities."""
    service = TestCaseService(db)
    filters = build_filters(
        project_id=project_id,
        module_id=module_id,
        epic_id=epic_id,
        user_story_id=user_story_id,
        status=test_status,
        test_type=test_type,
        priority=priority,
        automation_status=automation_status,
        search=search,
    )
    result = await service.list_test_cases(
        page=page,
        page_size=page_size,
        sort_by=sort_by,
        sort_order=sort_order,
        filters=filters if filters else None,
    )
    return success_response(data=result, message="Test cases retrieved successfully")


@router.post(
    "",
    response_model=None,
    status_code=status.HTTP_201_CREATED,
    summary="Create test case",
)
async def create_test_case(
    test_case_data: dict = Body(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """Create a new test case. Auto-generates a sequential ID (e.g., TC-00001)."""
    service = TestCaseService(db)
    test_case = await service.create_test_case(
        data=test_case_data, created_by=current_user["id"]
    )
    return success_response(data=test_case, message="Test case created successfully")


@router.get(
    "/{test_case_id}",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="Get test case by ID",
)
async def get_test_case(
    test_case_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """Get a specific test case by its UUID."""
    service = TestCaseService(db)
    test_case = await service.get_test_case(test_case_id=test_case_id)
    return success_response(data=test_case, message="Test case retrieved successfully")


@router.put(
    "/{test_case_id}",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="Update test case",
)
async def update_test_case(
    test_case_id: str,
    test_case_data: dict = Body(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """Update a test case. Creates a new version entry for audit trail."""
    service = TestCaseService(db)
    test_case = await service.update_test_case(
        test_case_id=test_case_id,
        data=test_case_data,
    )
    return success_response(data=test_case, message="Test case updated successfully")


@router.delete(
    "/{test_case_id}",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="Soft delete test case",
)
async def delete_test_case(
    test_case_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """Soft delete a test case."""
    service = TestCaseService(db)
    await service.delete_test_case(test_case_id=test_case_id)
    return success_response(message="Test case deleted successfully")


@router.post(
    "/{test_case_id}/clone",
    response_model=None,
    status_code=status.HTTP_201_CREATED,
    summary="Clone test case",
)
async def clone_test_case(
    test_case_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """Clone an existing test case, generating a new ID and resetting status."""
    service = TestCaseService(db)
    cloned = await service.clone_test_case(
        source_id=test_case_id, created_by=current_user["id"]
    )
    return success_response(data=cloned, message="Test case cloned successfully")


@router.get(
    "/{test_case_id}/versions",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="Get test case version history",
)
async def get_test_case_versions(
    test_case_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """Get the full version history for a test case."""
    service = TestCaseService(db)
    versions = await service.get_versions(test_case_id=test_case_id)
    return success_response(data=versions, message="Version history retrieved successfully")


@router.post(
    "/{test_case_id}/approve",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="Approve or reject test case",
)
async def approve_test_case(
    test_case_id: str,
    approval_data: dict = Body(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """Approve or reject a test case in the approval workflow.

    Expected approval_data: {"action": "approve" | "reject", "comment": "optional comment"}
    """
    service = TestCaseService(db)
    result = await service.process_approval(
        test_case_id=test_case_id,
        action=approval_data["action"],
        comment=approval_data.get("comment"),
        approved_by=current_user["id"],
    )
    return success_response(data=result, message="Approval processed successfully")


@router.post(
    "/bulk-upload",
    response_model=None,
    status_code=status.HTTP_201_CREATED,
    summary="Bulk upload test cases via Excel/CSV",
)
async def bulk_upload_test_cases(
    file: UploadFile = File(..., description="Excel or CSV file containing test cases"),
    project_id: str = Query(description="Target project ID"),
    epic_id: str | None = Query(default=None, description="Target epic ID"),
    user_story_id: str | None = Query(default=None, description="Target user story ID"),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    """Upload multiple test cases from an Excel or CSV file."""
    import csv as csv_mod
    import io

    service = TestCaseService(db)
    content = await file.read()
    filename = (file.filename or "").lower()

    rows: list[dict] = []
    if filename.endswith((".xlsx", ".xls")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        headers = [str(cell.value or "").strip().lower() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for h, v in zip(headers, row):
                if h and v is not None:
                    row_dict[h] = str(v).strip()
            if row_dict.get("title"):
                rows.append(row_dict)
        wb.close()
    else:
        text = content.decode("utf-8-sig")
        reader = csv_mod.DictReader(io.StringIO(text))
        for row in reader:
            row_dict = {k.strip().lower(): v.strip() for k, v in row.items() if k and v}
            if row_dict.get("title"):
                rows.append(row_dict)

    created = []
    errors = []
    for i, row_data in enumerate(rows, start=2):
        try:
            import json as json_mod
            row_data["project_id"] = project_id
            row_data["epic_id"] = epic_id
            row_data["user_story_id"] = user_story_id
            # Convert isAutomated boolean to automation_status
            is_automated = row_data.pop("isautomated", row_data.pop("is_automated", None))
            if is_automated is not None:
                row_data["automation_status"] = "Automated" if str(is_automated).lower() in ("true", "1", "yes") else "Manual"
            # Parse steps JSON string into a list
            if "steps" in row_data and isinstance(row_data["steps"], str):
                try:
                    row_data["steps"] = json_mod.loads(row_data["steps"])
                except (json_mod.JSONDecodeError, ValueError):
                    row_data["steps"] = None
            # Parse version to int
            if "version" in row_data and row_data["version"]:
                try:
                    row_data["version"] = int(float(row_data["version"]))
                except (ValueError, TypeError):
                    row_data.pop("version", None)
            tc = await service.create_test_case(row_data, created_by=current_user["id"])
            created.append({"row": i, "id": str(tc.id), "title": tc.title})
        except Exception as exc:
            errors.append({"row": i, "title": row_data.get("title", ""), "error": str(exc)})

    return success_response(
        data={"created": len(created), "failed": len(errors), "errors": errors},
        message=f"{len(created)} test case(s) created, {len(errors)} failed",
    )


# ── Bulk operations ────────────────────────────────────────────────────────
#
# Each endpoint accepts up to 500 IDs and returns a structured
# {succeeded, failed} report. Per-id failures (NotFound, illegal
# transition, validation) are non-fatal — the route always returns 200
# with the partial result so the UI can show progress.


@router.post(
    "/bulk-delete",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="Soft-delete multiple test cases",
)
async def bulk_delete_test_cases(
    payload: _BulkIds = Body(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    service = TestCaseService(db)
    result = await service.bulk_delete(payload.ids)
    return success_response(
        data=result,
        message=f"{len(result['succeeded'])} test case(s) deleted",
    )


@router.post(
    "/bulk-transition",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="Run the approval-workflow transition on multiple test cases",
)
async def bulk_transition_test_cases(
    payload: _BulkApprovalTransition = Body(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    service = TestCaseService(db)
    result = await service.bulk_transition_approval(payload.ids, payload.status)
    return success_response(
        data=result,
        message=f"{len(result['succeeded'])} test case(s) transitioned to '{payload.status}'",
    )


@router.post(
    "/bulk-update",
    response_model=None,
    status_code=status.HTTP_200_OK,
    summary="Bulk-update plain fields (priority, type, automation_status, assignee)",
)
async def bulk_update_test_cases(
    payload: _BulkUpdate = Body(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
) -> dict:
    has_change = (
        payload.priority is not None
        or payload.type_ is not None
        or payload.automation_status is not None
        or payload.assigned_to is not None
        or payload.unassign
    )
    if not has_change:
        raise ValidationError(
            "At least one of priority, type, automation_status, assigned_to, or unassign must be provided"
        )
    service = TestCaseService(db)
    result = await service.bulk_update(
        payload.ids,
        priority=payload.priority,
        type_=payload.type_,
        automation_status=payload.automation_status,
        assigned_to=payload.assigned_to,
        unassign=payload.unassign,
    )
    return success_response(
        data=result,
        message=f"{len(result['succeeded'])} test case(s) updated",
    )
